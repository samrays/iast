"""Convert Chapter 4 Markdown document into a beautifully formatted Word Document (.docx)."""

import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


def set_cell_background(cell, fill_hex):
    """Set background shading for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set internal cell padding."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_formatted_paragraph(doc, text, style='Normal', space_after=6, space_before=0, bold=False, italic=False, font_size=11, color_rgb=(50, 50, 50)):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15

    # Inline formatting parser for **bold** and *italic* and `code`
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\$.*?\$)', text)
    for token in tokens:
        if not token:
            continue
        run = p.add_run()
        if token.startswith('**') and token.endswith('**'):
            run.text = token[2:-2]
            run.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run.text = token[1:-1]
            run.italic = True
        elif token.startswith('`') and token.endswith('`'):
            run.text = token[1:-1]
            run.font.name = 'Consolas'
            run.font.size = Pt(font_size - 1)
            run.font.color.rgb = RGBColor(180, 40, 40)
        elif token.startswith('$') and token.endswith('$'):
            run.text = token[1:-1]
            run.font.name = 'Cambria Math'
            run.italic = True
        else:
            run.text = token
            run.bold = bold
            run.italic = italic

        if not (token.startswith('`') and token.endswith('`')):
            run.font.name = 'Calibri'
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(*color_rgb)

    return p


def convert_md_to_docx(md_filepath, docx_filepath):
    with open(md_filepath, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()

    # Set page margins (1 inch all around)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    in_code_block = False
    code_block_text = []
    in_table = False
    table_lines = []

    def flush_table(lines):
        if not lines:
            return
        rows_data = []
        for l in lines:
            if re.match(r'^\s*\|?\s*[-:]+\s*\|', l):
                continue
            parts = [c.strip() for c in l.strip().strip('|').split('|')]
            if parts and any(parts):
                rows_data.append(parts)

        if not rows_data:
            return

        cols_cnt = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=cols_cnt)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for r_idx, row in enumerate(rows_data):
            tr = table.rows[r_idx]
            is_header = (r_idx == 0)
            for c_idx, cell_value in enumerate(row):
                if c_idx < len(tr.cells):
                    cell = tr.cells[c_idx]
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                    
                    if is_header:
                        set_cell_background(cell, "1F4E78") # Navy header
                    elif r_idx % 2 == 1:
                        set_cell_background(cell, "F2F4F7") # Light zebra stripe
                    else:
                        set_cell_background(cell, "FFFFFF")

                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)

                    # Inline formatting
                    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', cell_value)
                    for token in tokens:
                        if not token:
                            continue
                        run = p.add_run()
                        if token.startswith('**') and token.endswith('**'):
                            run.text = token[2:-2]
                            run.bold = True
                        elif token.startswith('*') and token.endswith('*'):
                            run.text = token[1:-1]
                            run.italic = True
                        elif token.startswith('`') and token.endswith('`'):
                            run.text = token[1:-1]
                            run.font.name = 'Consolas'
                            run.font.size = Pt(9.5)
                        else:
                            run.text = token

                        run.font.name = 'Calibri'
                        run.font.size = Pt(9.5 if not is_header else 10)
                        if is_header:
                            run.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                        else:
                            run.font.color.rgb = RGBColor(40, 40, 40)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def flush_code_block(text_lines):
        if not text_lines:
            return
        code_text = "".join(text_lines)
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, "F4F5F7")
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.0)
        run.font.color.rgb = RGBColor(30, 40, 60)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    for line in lines:
        raw_line = line.rstrip('\n')
        stripped = raw_line.strip()

        # Handle Code blocks
        if stripped.startswith('```'):
            if in_code_block:
                flush_code_block(code_block_text)
                code_block_text = []
                in_code_block = False
            else:
                if in_table:
                    flush_table(table_lines)
                    table_lines = []
                    in_table = False
                in_code_block = True
            continue

        if in_code_block:
            code_block_text.append(raw_line + "\n")
            continue

        # Handle Tables
        if '|' in stripped and (stripped.startswith('|') or stripped.endswith('|')):
            if not in_table:
                in_table = True
            table_lines.append(stripped)
            continue
        else:
            if in_table:
                flush_table(table_lines)
                table_lines = []
                in_table = False

        if not stripped:
            continue

        # Handle Markdown Images ![alt](src)
        img_match = re.match(r'^\s*\!\[(.*?)\]\((.*?)\)\s*$', stripped)
        if img_match:
            alt_text, img_src = img_match.groups()
            if img_src.startswith("file:///"):
                img_src = img_src[8:]
            img_src = img_src.replace('/', os.sep)
            if os.path.exists(img_src):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run()
                run.add_picture(img_src, width=Inches(6.0))
                
                cp = doc.add_paragraph()
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cp.paragraph_format.space_before = Pt(2)
                cp.paragraph_format.space_after = Pt(14)
                crun = cp.add_run(alt_text)
                crun.font.name = 'Calibri'
                crun.font.size = Pt(9.5)
                crun.italic = True
                crun.font.color.rgb = RGBColor(100, 100, 100)
            continue

        # Headings
        if stripped.startswith('# '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(18)
            h.paragraph_format.space_after = Pt(12)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = h.add_run(stripped[2:])
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(22)
            run.font.color.rgb = RGBColor(31, 78, 120)
        elif stripped.startswith('## '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(8)
            run = h.add_run(stripped[3:])
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(31, 78, 120)
        elif stripped.startswith('### '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(6)
            run = h.add_run(stripped[4:])
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(46, 117, 182)
        elif stripped.startswith('#### '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            run = h.add_run(stripped[5:])
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(11.5)
            run.font.color.rgb = RGBColor(50, 50, 50)
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = add_formatted_paragraph(doc, stripped[2:], style='List Bullet', space_after=4, font_size=11)
        elif re.match(r'^\d+\.\s+', stripped):
            content = re.sub(r'^\d+\.\s+', '', stripped)
            p = add_formatted_paragraph(doc, content, style='List Number', space_after=4, font_size=11)
        elif stripped == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run('―' * 40)
            run.font.color.rgb = RGBColor(180, 180, 180)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            add_formatted_paragraph(doc, stripped, space_after=6, font_size=11)

    if in_table:
        flush_table(table_lines)

    doc.save(docx_filepath)
    print(f"Successfully generated Word Document at: {docx_filepath}")


if __name__ == "__main__":
    md_path = r"c:\Users\CP-1005\gravity\iast\iast\Chapter_4_Implementation_and_Results.md"
    docx_path = r"c:\Users\CP-1005\gravity\iast\iast\Chapter_4_Implementation_and_Results.docx"
    convert_md_to_docx(md_path, docx_path)
