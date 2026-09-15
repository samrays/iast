"""Synchronize Table Titles and Data Parity Across all DOCX Document Files."""

import os
import glob
import docx

def update_docx_files():
    print("Synchronizing DOCX documents for 100% data parity...")

    # 1. Update Chapter_5_Discussion.docx
    c5_path = "Chapter_5_Discussion.docx"
    if os.path.exists(c5_path):
        doc = docx.Document(c5_path)
        for idx, p in enumerate(list(doc.paragraphs)):
            if "5.2.2 What Vulnerability Types Were Most / Least Detected by DAST?" in p.text:
                if not any("Table 5.1" in p2.text for p2 in doc.paragraphs[max(0, idx-2):idx+5]):
                    p.insert_paragraph_before("### Table 5.1: DAST Vulnerability Detection Efficacy & Recall Breakdown by OWASP Category")
            if "5.3.2 What Vulnerability Types Were Most / Least Detected by IAST?" in p.text:
                if not any("Table 5.3" in p2.text for p2 in doc.paragraphs[max(0, idx-2):idx+5]):
                    p.insert_paragraph_before("### Table 5.3: IAST Vulnerability Detection Efficacy & Recall Breakdown by OWASP Category")
        doc.save(c5_path)
        print(f"Updated {c5_path}")

    # 2. Update Chapter_4_Implementation_and_Results.docx
    c4_path = "Chapter_4_Implementation_and_Results.docx"
    if os.path.exists(c4_path):
        doc = docx.Document(c4_path)
        for idx, p in enumerate(list(doc.paragraphs)):
            if "4.4.1 Precision, Recall, and F1-Score" in p.text:
                if not any("Table 4.3" in p2.text for p2 in doc.paragraphs[max(0, idx-2):idx+5]):
                    p.insert_paragraph_before("### Table 4.3: DAST (OWASP ZAP v2.15.0) Contingency Matrix")
            if "4.7.2 Ground Truth Vulnerability Mapping Matrix" in p.text:
                if not any("Table 4.10" in p2.text for p2 in doc.paragraphs[max(0, idx-2):idx+5]):
                    p.insert_paragraph_before("### Table 4.10: Ground-Truth Vulnerability Detection & Category Breakdown Matrix (DAST vs. IAST)")
        doc.save(c4_path)
        print(f"Updated {c4_path}")

    # 3. Update Appendices.docx and Chapter_6_Appendices.docx
    for app_path in ["Appendices.docx", "Chapter_6_Appendices.docx"]:
        if os.path.exists(app_path):
            doc = docx.Document(app_path)
            for table in doc.tables:
                if len(table.rows) >= 11:
                    row0_text = " ".join([c.text.strip() for c in table.rows[0].cells])
                    if "OWASP Vulnerability Category" in row0_text and "Detected (TP)" in row0_text:
                        expected_data = [
                            ("4", "80.00% (4/5)"),
                            ("4", "80.00% (4/5)"),
                            ("3", "75.00% (3/4)"),
                            ("4", "80.00% (4/5)"),
                            ("5", "83.33% (5/6)"),
                            ("5", "71.43% (5/7)"),
                            ("4", "80.00% (4/5)"),
                            ("5", "83.33% (5/6)"),
                            ("4", "80.00% (4/5)"),
                            ("5", "71.43% (5/7)"),
                        ]
                        for idx, (tp_val, prec_val) in enumerate(expected_data, start=1):
                            if idx < len(table.rows):
                                r_cells = table.rows[idx].cells
                                if len(r_cells) >= 7:
                                    r_cells[3].text = tp_val
                                    r_cells[6].text = prec_val
                        sum_cells = table.rows[11].cells
                        sum_cells[3].text = "43"
                        sum_cells[6].text = "78.18% (43/55)"
            doc.save(app_path)
            print(f"Updated {app_path}")

    print("All DOCX files successfully synchronized.")

if __name__ == "__main__":
    update_docx_files()
